"""Render the illustrated Markdown manual to PDF or DOCX; retain original screenshots.

Requires reportlab, python-docx and a Chinese font. No business data is read.
Example: python docs/tools/build-operation-manual.py --font C:/Windows/Fonts/simhei.ttf
"""

import argparse
import html
import re
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, KeepTogether, PageBreak, PageTemplate,
    Paragraph, Spacer, Table, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/operation-manual.md"
BLUE = colors.HexColor("#153D5B")
INK = colors.HexColor("#263A4D")
MUTED = colors.HexColor("#61758A")
PALE = colors.HexColor("#EDF4FB")
WIDTH, HEIGHT = A4
MARGIN = 42
CONTENT_WIDTH = WIDTH - 2 * MARGIN


def markup(text):
    """Minimal Markdown inline formatting; never interpret raw HTML as instructions."""
    text = html.escape(text, quote=False)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`([^`]+)`", r'<font color="#153D5B">\1</font>', text)

    def link(match):
        label, target = match.groups()
        if target.startswith("http"):
            url = target
        else:
            path = (SOURCE.parent / html.unescape(target)).resolve().relative_to(ROOT)
            url = "https://github.com/Saikk1/meeting-list-ai/blob/main/" + path.as_posix()
        return f'<link href="{html.escape(url, quote=True)}" color="#2167B0">{label}</link>'

    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, text)


class ManualDoc(BaseDocTemplate):
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and getattr(flowable, "chapter", False):
            label = flowable.getPlainText()
            key = "chapter-" + label.split(".", 1)[0]
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(label, key, level=0, closed=False)
            self.notify("TOCEntry", (0, label, self.page, key))


def decorate(canvas, doc):
    canvas.saveState()
    if doc.page == 1:
        canvas.setFillColor(BLUE)
        canvas.rect(0, HEIGHT - 18, WIDTH, 18, fill=1, stroke=0)
    else:
        canvas.setFont("ManualCN", 8.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(MARGIN, HEIGHT - 28, "会务智办  |  图文操作手册")
        canvas.drawRightString(WIDTH - MARGIN, HEIGHT - 28, "3.1.4")
        canvas.setStrokeColor(colors.HexColor("#D9E4EE"))
        canvas.line(MARGIN, HEIGHT - 35, WIDTH - MARGIN, HEIGHT - 35)
    canvas.setFillColor(MUTED)
    canvas.setFont("ManualCN", 8)
    canvas.drawString(MARGIN, 23, "演示截图不含真实业务数据  ·  2026-10-27")
    canvas.drawRightString(WIDTH - MARGIN, 23, str(doc.page))
    canvas.restoreState()


def build(output, font):
    pdfmetrics.registerFont(TTFont("ManualCN", font))
    pdfmetrics.registerFontFamily("ManualCN", normal="ManualCN", bold="ManualCN")
    body = ParagraphStyle("Body", fontName="ManualCN", fontSize=10.5,
                          leading=17, textColor=INK, wordWrap="CJK", spaceAfter=9)
    title = ParagraphStyle("Chapter", parent=body, fontSize=18, leading=26,
                           textColor=BLUE, spaceAfter=14)
    note = ParagraphStyle("Note", parent=body, fontSize=9.5, leading=16,
                          backColor=PALE, borderPadding=9, spaceBefore=6, spaceAfter=12)
    caption = ParagraphStyle("Caption", parent=body, fontSize=8.5, leading=13,
                             textColor=MUTED, alignment=1, spaceAfter=12)
    cell = ParagraphStyle("Cell", parent=body, fontSize=9, leading=15, spaceAfter=0)
    toc_style = ParagraphStyle("Contents", parent=body, fontSize=10, leading=16,
                              leftIndent=0, firstLineIndent=0, spaceAfter=1)

    raw = SOURCE.read_text(encoding="utf-8")
    parts = re.split(r"^## (.+)$", raw, flags=re.M)
    doc = ManualDoc(str(output), pagesize=A4, leftMargin=MARGIN, rightMargin=MARGIN,
                    topMargin=51, bottomMargin=43, title="会务智办操作手册 3.1.4",
                    author="会务智办", pageCompression=1)
    doc.addPageTemplates([PageTemplate(id="Manual", frames=Frame(
        MARGIN, 43, CONTENT_WIDTH, HEIGHT - 94, leftPadding=0, rightPadding=0,
        topPadding=0, bottomPadding=0), onPage=decorate)])
    story = [Spacer(1, 72)]
    story.append(Paragraph("会务智办", ParagraphStyle("Brand", parent=title,
                                                     fontSize=24, leading=34)))
    story.append(Paragraph("操作手册", ParagraphStyle("Cover", parent=title,
                                                    fontSize=38, leading=50, spaceAfter=25)))
    story.append(Paragraph("从新建会议到名单、车辆、座次与会务跟进", body))
    story.append(Spacer(1, 18))
    intro = parts[0].split("\n\n")
    for paragraph in intro[1:]:
        if paragraph.strip():
            story.append(Paragraph(markup(paragraph.replace("\n", " ")), body))
    story.extend([Spacer(1, 28), Paragraph(
        "保存会议与保留草稿不同。人员、车辆、正式名单、座次分别管理；按本次会议需要启用。", note),
        PageBreak(), Paragraph("目录", title),
        Paragraph("点击章节或页码跳转。截图保留原始分辨率，可放大查看。", body)])
    toc = TableOfContents()
    toc.levelStyles = [toc_style]
    toc.dotsMinLevel = 0
    story.append(toc)

    for index in range(1, len(parts), 2):
        heading, content = parts[index], parts[index + 1]
        story.append(PageBreak())
        chapter = Paragraph(markup(heading), title)
        chapter.chapter = True
        story.append(chapter)
        lines = content.strip().splitlines()
        pending = []

        def flush():
            if pending:
                story.append(Paragraph(markup(" ".join(pending)), body))
                pending.clear()

        pos = 0
        while pos < len(lines):
            line = lines[pos].strip()
            image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if not line:
                flush()
            elif image_match:
                flush()
                alt, filename = image_match.groups()
                pic = Image(str(SOURCE.parent / filename))
                ratio = min(CONTENT_WIDTH / pic.imageWidth, 285 / pic.imageHeight)
                pic.drawWidth, pic.drawHeight = pic.imageWidth * ratio, pic.imageHeight * ratio
                story.append(KeepTogether([pic, Spacer(1, 5), Paragraph(markup(alt), caption)]))
            elif line.startswith("|"):
                flush()
                table_lines = []
                while pos < len(lines) and lines[pos].strip().startswith("|"):
                    table_lines.append(lines[pos].strip())
                    pos += 1
                rows = []
                for row in table_lines:
                    if re.fullmatch(r"[| :\-]+", row):
                        continue
                    values = row.strip("|").split("|")
                    rows.append([Paragraph(markup(value.strip()), cell) for value in values])
                widths = [CONTENT_WIDTH * .34, CONTENT_WIDTH * .66] if len(rows[0]) == 2 else None
                table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), PALE),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("LINEBELOW", (0, 0), (-1, 0), .6, colors.HexColor("#CDDCEB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]))
                story.extend([table, Spacer(1, 12)])
                continue
            elif line.startswith("> "):
                flush()
                story.append(Paragraph(markup(line[2:]), note))
            elif re.match(r"^\d+\. ", line):
                flush()
                story.append(Paragraph(markup(line), body))
            else:
                pending.append(line)
            pos += 1
        flush()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.multiBuild(story)
    print(f"PDF created: {output}")


def build_word(output):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(1)
    section.footer_distance = Inches(.35)
    for name, size in [("Normal", 11), ("Title", 30), ("Heading 1", 18), ("Caption", 9)]:
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.25
    document.styles["Title"].paragraph_format.space_after = Pt(22)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    document.styles["Caption"].paragraph_format.space_after = Pt(10)
    # The runtime's base Word template may carry a blue title paragraph border.
    # Remove inherited decoration so the manual title is plain black typography.
    for element in document.styles.element.iter():
        if element.tag == qn("w:pBdr"):
            element.getparent().remove(element)
    properties = document.core_properties
    properties.title = "会务智办操作手册"
    properties.subject = "3.1.4 会务与座次工作台 图文操作步骤"
    properties.author = "会务智办"
    properties.last_modified_by = "会务智办"

    def text(paragraph, value):
        # Keep readable labels in Word; online and download entry URLs are in Markdown.
        value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", value)
        for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", value):
            run = paragraph.add_run(token.strip("*") if token.startswith("**") else token.strip("`"))
            if token.startswith("**"):
                run.bold = True

    parts = re.split(r"^## (.+)$", SOURCE.read_text(encoding="utf-8"), flags=re.M)
    document.add_paragraph("会务智办操作手册", "Title")
    document.add_paragraph("会议安排  人员名单  车辆登记  座次编排  会务跟进")
    document.add_paragraph("3.1.4    更新日期 2026-10-27")
    for paragraph in parts[0].split("\n\n")[2:]:
        if paragraph.strip():
            text(document.add_paragraph(), paragraph.replace("\n", " "))
    contents_heading = document.add_paragraph("目录", "Heading 1")
    contents_heading.paragraph_format.page_break_before = True
    for index in range(1, len(parts), 2):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1
        anchor = OxmlElement("w:hyperlink")
        anchor.set(qn("w:anchor"), "chapter_" + str(index))
        run = OxmlElement("w:r")
        content = OxmlElement("w:t")
        content.text = re.sub(r"[.：:]+", " ", parts[index])
        run.append(content)
        anchor.append(run)
        p._p.append(anchor)

    for index in range(1, len(parts), 2):
        heading = document.add_paragraph(re.sub(r"[.：:]+", " ", parts[index]), "Heading 1")
        heading.paragraph_format.page_break_before = True
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), "chapter_" + str(index))
        end.set(qn("w:id"), str(index))
        heading._p.insert(0, start)
        heading._p.append(end)
        lines = parts[index + 1].strip().splitlines()
        pos = 0
        while pos < len(lines):
            line = lines[pos].strip()
            match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if not line:
                pos += 1
                continue
            if match:
                alt, filename = match.groups()
                p = document.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shape = p.add_run().add_picture(str(SOURCE.parent / filename), width=Inches(6.5))
                # Keep the expanded attendee instructions together with their screenshot.
                image_height = Inches(3.1 if index == 9 or filename.endswith(("room-table-config.png", "manual-ux/quick-navigation.png")) else 3.7)
                if shape.height > image_height:
                    shape.width = int(shape.width * image_height / shape.height)
                    shape.height = image_height
                shape._inline.docPr.set("descr", alt)
                caption_p = document.add_paragraph(alt, "Caption")
                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith("|"):
                rows = []
                while pos < len(lines) and lines[pos].strip().startswith("|"):
                    row = lines[pos].strip()
                    if not re.fullmatch(r"[| :\-]+", row):
                        rows.append([v.strip() for v in row.strip("|").split("|")])
                    pos += 1
                table = document.add_table(rows=0, cols=len(rows[0]))
                table.autofit = False
                table.columns[0].width, table.columns[1].width = Inches(2.2), Inches(4.3)
                borders = OxmlElement("w:tblBorders")
                for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                    border = OxmlElement("w:" + edge)
                    for key, value in [("val", "single"), ("sz", "6"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for number, values in enumerate(rows):
                    cells = table.add_row().cells
                    for column, (value, cell) in enumerate(zip(values, cells)):
                        cell.width = Inches(2.2 if column == 0 else 4.3)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        p = cell.paragraphs[0]
                        text(p, value)
                        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.1
                        for run in p.runs:
                            run.font.size = Pt(10)
                            run.bold = number == 0
                        margins = OxmlElement("w:tcMar")
                        for edge in ["top", "left", "bottom", "right"]:
                            margin = OxmlElement("w:" + edge)
                            margin.set(qn("w:w"), "90")
                            margin.set(qn("w:type"), "dxa")
                            margins.append(margin)
                        cell._tc.get_or_add_tcPr().append(margins)
                        if number == 0:
                            shade = OxmlElement("w:shd")
                            shade.set(qn("w:fill"), "EAF1F7")
                            cell._tc.get_or_add_tcPr().append(shade)
                    if number == 0:
                        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                document.add_paragraph()
                continue
            else:
                p = document.add_paragraph()
                text(p, line[2:] if line.startswith("> ") else line)
                if index == 9:
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.line_spacing = 1.15
                elif index >= 49:
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.size = Pt(10.5)
            pos += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("会务智办操作手册  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Word created: {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--font", default="C:/Windows/Fonts/simhei.ttf")
    parser.add_argument("--format", choices=["pdf", "docx"], default="pdf")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.format == "docx":
        build_word(args.output or ROOT / "docs/downloads/会务智办操作手册.docx")
    else:
        build(args.output or ROOT / "output/pdf/会务智办操作手册.pdf", args.font)
